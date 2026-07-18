import re
import os
from docx import Document
PLACEHOLDER_PATTERN = r"\{\{(.*?)\}\}"  
def extract_placeholders(template_text: str):
    """
    Extracts placeholders from a template like:
    'Hello {{name}}, your department is {{department}}'
    Returns: ['name', 'department']
    """
    return re.findall(PLACEHOLDER_PATTERN, template_text)
def replace_placeholders(template_text: str, values: dict):
    """
    Replaces placeholders with the provided dictionary values.
    """
    def replacer(match):
        key = match.group(1)
        return str(values.get(key, f"<missing:{key}>"))
    return re.sub(PLACEHOLDER_PATTERN, replacer, template_text)
def load_template_text(file_path):
    """
    Loads text from a .txt or .docx template file.
    Returns the text as a string.
    """
    if not os.path.exists(file_path):
        return ""

    if file_path.lower().endswith(".txt"):
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    elif file_path.lower().endswith(".docx"):
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return "\n".join(full_text)
    else:
        raise ValueError("Unsupported file type: " + file_path)


try:
    import spacy
    nlp = spacy.load("en_core_web_sm")

    def analyze_text(text: str):
        """
        Basic NLP analysis using spaCy:
        - Named entity recognition
        - Noun phrase extraction
        """
        doc = nlp(text)
        return {
            "entities": [(ent.text, ent.label_) for ent in doc.ents],
            "noun_phrases": [chunk.text for chunk in doc.noun_chunks]
        }

except Exception:
    
    def analyze_text(text: str):
        return {
            "entities": [],
            "noun_phrases": []
        }



if __name__ == "__main__":
    sample = "Hello {{name}}, your department is {{department}}"
    print("Placeholders:", extract_placeholders(sample))
    print("Replaced:", replace_placeholders(sample, {"name": "Srihari", "department": "CSE"}))